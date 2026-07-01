#!/usr/bin/env python3
"""
從公開的 Google Sheet（請款資料）抓取「日記帳」，聚合成：
  - 各專案 × 各月份 的「收入」與「淨額（收入-成本）」
  - 每月 預收 / 應收 / 應付 帳戶進出
輸出 data/projects_monthly.json，供中台 index.html 使用。

用法：
  python3 scripts/build_data.py            # 線上抓最新公開 Sheet
  python3 scripts/build_data.py book.xlsx  # 用本地 xlsx
之後重新部署 Vercel 即可同步更新。
"""
import sys, json, os, datetime, urllib.request
import openpyxl

SHEET_ID = "17NmFtx8XKjB5_4XR9N1FHJPZWWB2IAZW810bTVdbmYM"
XLSX_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
OUT = os.path.join(os.path.dirname(__file__), "..", "data", "projects_monthly.json")

def load_book(path=None):
    if path:
        return openpyxl.load_workbook(path, data_only=True)
    tmp = "/tmp/_journal.xlsx"
    urllib.request.urlretrieve(XLSX_URL, tmp)
    return openpyxl.load_workbook(tmp, data_only=True)

def find_sheet(wb, key):
    for ws in wb.worksheets:
        if key in ws.title:
            return ws
    raise KeyError(key)

def main():
    wb = load_book(sys.argv[1] if len(sys.argv) > 1 else None)
    cs = find_sheet(wb, "分類")
    cat2big = {}
    for r in range(6, cs.max_row + 1):
        n = cs.cell(r, 2).value
        if n:
            cat2big[str(n).strip()] = (cs.cell(r, 3).value or "").strip()

    ws = find_sheet(wb, "日記帳")
    rev, cost = {}, {}
    acct = {"預計收入": [0]*12, "應收帳款": [0]*12, "應付帳款": [0]*12}
    ar_book, ar_collected = [], 0.0   # 應收帳款：認列明細 / 已收回總額
    billing = []                      # 每筆收入的實際收款狀態（供時間軸自動對帳）
    collected_descs = []              # 已收款分錄的描述（把應收/預收轉入銀行）
    for r in range(4, ws.max_row + 1):
        d = ws.cell(r, 2).value
        desc = ws.cell(r, 3).value
        amt = ws.cell(r, 5).value
        cat = ws.cell(r, 4).value
        outacc = ws.cell(r, 6).value
        inacc = ws.cell(r, 7).value
        proj = ws.cell(r, 8).value
        if not isinstance(d, datetime.datetime) or d.year != 2026:
            continue
        m = d.month - 1
        amt = float(amt) if amt not in (None, "") else 0.0
        big = cat2big.get(str(cat).strip(), "") if cat else ""
        pname = str(proj).strip() if proj else "（未標專案）"
        # 收款分錄：把應收帳款/預計收入 轉入 真實帳戶（銀行）＝該筆已收現
        if outacc in ("應收帳款", "預計收入") and inacc and inacc not in ("應收帳款", "預計收入"):
            collected_descs.append(str(desc or "").strip())
        # 收入收款明細（含未分類但轉入應收/預收的認列）：供時間軸自動對帳
        if amt > 0 and (big == "收入" or inacc in ("應收帳款", "預計收入")):
            # 轉入應收帳款=應收未收；轉入預計收入=預計；其餘(銀行等)=已收
            status = "應收" if inacc == "應收帳款" else ("預計" if inacc == "預計收入" else "已收")
            billing.append({"proj": pname, "desc": str(desc or "").strip(),
                            "m": d.month, "amt": round(amt), "status": status})
        if big == "收入":
            rev.setdefault(pname, [0]*12)[m] += amt
        elif big == "成本":
            cost.setdefault(pname, [0]*12)[m] += amt
        # 帳戶每月進出（資產：轉入為+，轉出為−）
        for a in acct:
            if inacc == a:
                acct[a][m] += amt
            if outacc == a:
                acct[a][m] -= amt
        # 應收帳款明細：認列（轉入應收）與收回（轉出應收）
        if inacc == "應收帳款" and amt > 0:
            ar_book.append([d.date().isoformat(), pname, str(desc or "").strip(), amt])
        if outacc == "應收帳款":
            ar_collected += amt

    # 應收未收＝已認列到應收帳款、且「沒有」對應收款分錄(描述相符)的款
    ar_book.sort(key=lambda x: x[0])
    ar_open = []
    for dt, pname, desc, amt in ar_book:
        if desc and any(desc in cd for cd in collected_descs):
            continue  # 已收款，沖銷
        ar_open.append({"date": dt, "name": pname, "desc": desc, "amt": round(amt)})

    projects = sorted(set(rev) | set(cost),
                      key=lambda p: -sum(rev.get(p, [0]*12)))
    rows = []
    for p in projects:
        rv = [round(x) for x in rev.get(p, [0]*12)]
        ct = [round(x) for x in cost.get(p, [0]*12)]
        net = [rv[i] + ct[i] for i in range(12)]
        rows.append({"name": p, "rev": rv, "net": net})

    data = {
        "updated": datetime.date.today().isoformat(),
        "months": ["1月","2月","3月","4月","5月","6月","7月","8月","9月","10月","11月","12月"],
        "projects": rows,
        "accounts": {k: [round(x) for x in v] for k, v in acct.items()},
        "receivables_open": ar_open,
        "receivables_open_total": round(sum(x["amt"] for x in ar_open)),
        "billing": billing,
    }
    # 已認列的款若之後有對應收款分錄（描述相符）＝改標已收
    for b in billing:
        if b["status"] in ("應收", "預計") and b["desc"]:
            if any(b["desc"] in cd for cd in collected_descs):
                b["status"] = "已收"
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=0)
    tot = [sum(r["rev"][i] for r in rows) for i in range(12)]
    print("✓ 寫入", OUT)
    print("  專案數:", len(rows), " 全年收入:", f"{sum(tot):,}")

if __name__ == "__main__":
    main()
